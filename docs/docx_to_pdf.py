"""Convert DOCX report to PDF — DBIMSCA format, no table borders."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

DOCX = Path(r"C:\Users\jayat.ROHITH\Downloads\AthleteForge_Project_Report_50Pages.docx")
PDF = Path(r"C:\Users\jayat.ROHITH\Downloads\AthleteForge_Project_Report_50Pages.pdf")
BACKUP = Path(__file__).parent / "AthleteForge_Project_Report_50Pages.pdf"
FOOTER = "ATHLETEFORGE"


def esc(t: str) -> str:
    return (
        t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        .replace("\u2014", "-").replace("\u2013", "-")
        .replace("\u2018", "'").replace("\u2019", "'")
        .replace("\u201c", '"').replace("\u201d", '"')
    )


def styles():
    return {
        "chapter": ParagraphStyle("chapter", fontName="Times-Bold", fontSize=14, spaceBefore=16, spaceAfter=10, leading=18),
        "section": ParagraphStyle("section", fontName="Times-Bold", fontSize=12, spaceBefore=14, spaceAfter=8, leading=16),
        "body": ParagraphStyle("body", fontName="Times-Roman", fontSize=11, leading=16, alignment=TA_JUSTIFY, spaceAfter=8),
        "center": ParagraphStyle("center", fontName="Times-Roman", fontSize=11, leading=14, alignment=TA_CENTER, spaceAfter=6),
        "right": ParagraphStyle("right", fontName="Times-Roman", fontSize=11, alignment=TA_RIGHT, spaceAfter=6),
        "footer": ParagraphStyle("footer", fontName="Times-Roman", fontSize=9, alignment=TA_CENTER, textColor=colors.black),
        "table": ParagraphStyle("table", fontName="Times-Roman", fontSize=10, leading=13, spaceAfter=4),
        "thead": ParagraphStyle("thead", fontName="Times-Bold", fontSize=10, leading=13, spaceAfter=4),
    }


def is_chapter(text: str) -> bool:
    t = text.strip().upper()
    return t.startswith("CHAPTER ") or t in ("INTRODUCTION", "SYSTEM ANALYSIS", "SYSTEM DESIGN",
                                              "SYSTEM TESTING", "SYSTEM IMPLEMENTATION",
                                              "FUTURE SCOPE", "CONCLUSION", "REFERENCES",
                                              "ACKNOWLEDGEMENT", "TABLE OF CONTENTS")


def is_section(text: str) -> bool:
    t = text.strip()
    return len(t) >= 3 and t[0].isdigit() and "." in t[:4]


def page_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Times-Roman", 9)
    canvas.drawCentredString(A4[0] / 2, 1.0 * cm, f"{FOOTER}  {doc.page}  DBIMSCA")
    canvas.restoreState()


def add_table(story, table, st):
    rows = []
    for r_i, row in enumerate(table.rows):
        cells = [esc(c.text.strip().replace("\n", " ")) for c in row.cells]
        style = st["thead"] if r_i == 0 else st["table"]
        rows.append([Paragraph(c if c else " ", style) for c in cells])
    if not rows:
        return
    n = len(rows[0])
    w = 17 * cm / n
    t = Table(rows, colWidths=[w] * n)
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(Spacer(1, 6))
    story.append(t)
    story.append(Spacer(1, 8))


def iter_blocks(doc):
    """Yield paragraphs and tables in document order."""
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.text.paragraph import Paragraph as DocxP
    from docx.table import Table as DocxT
    for child in doc.element.body:
        if isinstance(child, CT_P):
            yield ("p", DocxP(child, doc))
        elif isinstance(child, CT_Tbl):
            yield ("t", DocxT(child, doc))


def convert():
    if not DOCX.exists():
        raise FileNotFoundError(f"Run: node generate-report.js first. Missing {DOCX}")

    doc = Document(str(DOCX))
    st = styles()
    story = []
    prev_chapter = False

    for kind, block in iter_blocks(doc):
        if kind == "t":
            add_table(story, block, st)
            continue

        text = block.text.strip()
        if not text:
            story.append(Spacer(1, 4))
            continue

        align = block.alignment
        if is_chapter(text):
            if prev_chapter and "CHAPTER" in text.upper():
                story.append(PageBreak())
            story.append(Paragraph(esc(text), st["chapter"]))
            prev_chapter = True
        elif is_section(text):
            story.append(Paragraph(esc(text), st["section"]))
        elif align == 1:  # CENTER
            story.append(Paragraph(esc(text), st["center"]))
        elif align == 2:  # RIGHT
            story.append(Paragraph(esc(text), st["right"]))
        else:
            story.append(Paragraph(esc(text), st["body"]))

    pdf = SimpleDocTemplate(
        str(PDF), pagesize=A4,
        leftMargin=2.5 * cm, rightMargin=2.5 * cm,
        topMargin=2.5 * cm, bottomMargin=2.2 * cm,
        title="AthleteForge Project Report",
        author="ROHITH GOWDA V & PRAKRUTHI R",
    )
    pdf.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    BACKUP.write_bytes(PDF.read_bytes())

    from pypdf import PdfReader
    pages = len(PdfReader(str(PDF)).pages)
    print(f"PDF: {PDF}")
    print(f"Pages: {pages}")
    print(f"Size: {PDF.stat().st_size / 1024:.1f} KB")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    convert()